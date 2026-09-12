import hashlib
import json
import os
import subprocess
import sys
import tempfile
import time
import unicodedata
from datetime import date, datetime, time as datetime_time
from collections.abc import Iterator
from contextlib import closing
from functools import lru_cache
from pathlib import Path


PROCESSOR_VERSION = "1.0.0"
PROCESSOR_IDS = {
    "cpu": "centaeris.document.cpu",
    "gpu:0": "centaeris.document.cuda.gpu0",
}
DET_MODEL = "PP-OCRv6_small_det"
REC_MODEL = "PP-OCRv6_small_rec"
RENDER_DPI = 220
MAX_INPUT_BYTES = 64 * 1024 * 1024
MAX_RENDERED_PIXELS_PER_PAGE = 16_000_000
MAX_OUTPUT_BYTES = 256 * 1024 * 1024
# Leave room for the existing 64 MiB commit metadata envelope.
MAX_MANIFEST_BYTES = 64 * 1024 * 1024 - 64 * 1024
MAX_PAGE_TEXT_BYTES = 4 * 1024 * 1024
MAX_WORKBOOK_PREVIEW_BYTES = 16 * 1024 * 1024
MAX_WORKBOOK_PREVIEW_CELLS = 250_000
MAX_WORKBOOK_PREVIEW_GRID_CELLS = 100_000
MAX_WORKBOOK_PREVIEW_SHEETS = 100
MODEL_MANIFEST = Path(os.environ.get("CENTAERIS_MODEL_MANIFEST", "/opt/centaeris/models/manifest.json"))
TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".json",
    ".jsonl",
    ".xml",
    ".yaml",
    ".yml",
    ".html",
    ".htm",
    ".css",
    ".js",
    ".mjs",
    ".ts",
    ".tsx",
    ".jsx",
    ".py",
    ".rs",
    ".toml",
}
OFFICE_SUFFIXES = {
    ".doc",
    ".docx",
    ".ppt",
    ".pptx",
    ".xls",
    ".xlsx",
    ".odt",
    ".ods",
    ".odp",
}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
MODEL_PAYLOAD_FILES = ("inference.json", "inference.pdiparams", "inference.yml")


class ProcessingError(RuntimeError):
    pass


def sha256_bytes(value: bytes) -> str:
    return f"sha256:{hashlib.sha256(value).hexdigest()}"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        while chunk := source.read(1024 * 1024):
            digest.update(chunk)
    return f"sha256:{digest.hexdigest()}"


def hash_model_payload(root: Path) -> str:
    digest = hashlib.sha256()
    files = [root / name for name in MODEL_PAYLOAD_FILES]
    if any(not path.is_file() for path in files):
        raise ProcessingError(f"model payload is incomplete: {root}")
    for path in files:
        relative = path.name.encode()
        digest.update(len(relative).to_bytes(4, "big"))
        digest.update(relative)
        with path.open("rb") as source:
            while chunk := source.read(1024 * 1024):
                digest.update(chunk)
    return f"sha256:{digest.hexdigest()}"


def fetch_models() -> None:
    from paddlex.inference.utils.official_models import official_models

    models = {}
    for name in (DET_MODEL, REC_MODEL):
        path = Path(official_models[name]).resolve()
        models[name] = {"path": str(path), "sha256": hash_model_payload(path)}
    MODEL_MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    MODEL_MANIFEST.write_text(
        json.dumps({"schema": "knowledge.model_manifest.v1", "models": models}, sort_keys=True),
        encoding="utf-8",
    )


@lru_cache(maxsize=1)
def load_model_manifest() -> dict:
    try:
        manifest = json.loads(MODEL_MANIFEST.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise ProcessingError("OCR model manifest is unavailable") from error
    if manifest.get("schema") != "knowledge.model_manifest.v1" or set(
        manifest.get("models", {})
    ) != {DET_MODEL, REC_MODEL}:
        raise ProcessingError("OCR model manifest is invalid")
    for name, item in manifest["models"].items():
        path = Path(item.get("path", ""))
        if not path.is_dir() or hash_model_payload(path) != item.get("sha256"):
            raise ProcessingError(f"OCR model digest mismatch: {name}")
    return manifest


def specification() -> dict:
    device = processor_device()
    require_device(device)
    manifest = load_model_manifest()
    return {
        "schema": "knowledge.processor_spec.v1",
        "processorId": PROCESSOR_IDS[device],
        "processorVersion": PROCESSOR_VERSION,
        "modelDigests": {
            name: manifest["models"][name]["sha256"] for name in (DET_MODEL, REC_MODEL)
        },
        "options": {
            "renderDpi": RENDER_DPI,
            "maxInputBytes": MAX_INPUT_BYTES,
            "maxRenderedPixelsPerPage": MAX_RENDERED_PIXELS_PER_PAGE,
            "maxOutputBytes": MAX_OUTPUT_BYTES,
        },
    }


def processor_device() -> str:
    value = os.environ.get("CENTAERIS_PROCESSOR_DEVICE", "")
    if value not in PROCESSOR_IDS:
        raise ProcessingError(
            f"CENTAERIS_PROCESSOR_DEVICE must be exactly cpu or gpu:0, got {value!r}"
        )
    return value


def require_device(device: str) -> None:
    if device == "cpu":
        return
    try:
        import paddle

        available = paddle.device.is_compiled_with_cuda()
        device_count = paddle.device.cuda.device_count()
    except Exception as error:
        raise ProcessingError(f"gpu:0 capability check failed: {error}") from error
    if not available or device_count != 1:
        raise ProcessingError(
            f"gpu:0 requires one CUDA-visible device, found {device_count}"
        )


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFC", value).replace("\r\n", "\n").replace("\r", "\n")
    return "".join(
        character
        for character in value
        if character in "\n\t" or not unicodedata.category(character).startswith("C")
    ).strip()


def native_text_is_usable(text: str) -> bool:
    nonspace = [character for character in text if not character.isspace()]
    if len(nonspace) < 16:
        return False
    replacements = sum(character == "\ufffd" for character in nonspace)
    printable = sum(character.isprintable() for character in nonspace)
    return replacements / len(nonspace) <= 0.02 and printable / len(nonspace) >= 0.95


def page_text(page: int, route: str, width: int, height: int, text: str, spans: list) -> dict:
    encoded = text.encode("utf-8")
    if len(encoded) > MAX_PAGE_TEXT_BYTES:
        raise ProcessingError(f"page {page} text exceeds the byte limit")
    return {
        "schema": "knowledge.page_text.v1",
        "page": page,
        "route": route,
        "widthMillipoints": max(1, width),
        "heightMillipoints": max(1, height),
        "text": text,
        "textSha256": sha256_bytes(encoded),
        "spans": spans,
    }


def empty_page(page: int, width: int, height: int) -> dict:
    return page_text(page, "empty", width, height, "", [])


def image_is_blank(image) -> bool:
    with image.convert("L") as grayscale:
        histogram = grayscale.histogram()
        marked = sum(histogram[:245])
        return marked / max(1, grayscale.width * grayscale.height) < 0.001


def _ocr_payload(result) -> dict:
    value = getattr(result, "json", None)
    value = value() if callable(value) else value
    if isinstance(value, str):
        value = json.loads(value)
    if not isinstance(value, dict):
        try:
            value = dict(result)
        except (TypeError, ValueError) as error:
            raise ProcessingError("PP-OCR result is not an object") from error
    value = value.get("res", value)
    if not isinstance(value, dict):
        raise ProcessingError("PP-OCR result payload is invalid")
    return value


def ocr_image(path: Path, page: int, width_millipoints: int, height_millipoints: int) -> dict:
    pipeline = ocr_pipeline()
    try:
        results = list(pipeline.predict(str(path)))
    except Exception as error:
        raise ProcessingError("PP-OCR inference failed") from error
    if len(results) != 1:
        raise ProcessingError("PP-OCR must return exactly one image result")
    payload = _ocr_payload(results[0])
    texts = list(payload.get("rec_texts", []))
    scores = list(payload.get("rec_scores", []))
    boxes = list(payload.get("rec_boxes", []))
    if not (len(texts) == len(scores) == len(boxes)):
        raise ProcessingError("PP-OCR result arrays have mismatched lengths")
    from PIL import Image

    with Image.open(path) as image:
        width, height = image.size
    spans = []
    for raw_text, raw_score, raw_box in zip(texts, scores, boxes, strict=True):
        text = normalize_text(str(raw_text))
        if not text:
            continue
        coordinates = [float(value) for value in raw_box]
        if len(coordinates) != 4:
            raise ProcessingError("PP-OCR box must contain four coordinates")
        left, top, right, bottom = coordinates
        bbox = [
            max(0, min(9_999, round(left * 10_000 / width))),
            max(0, min(9_999, round(top * 10_000 / height))),
            max(1, min(10_000, round(right * 10_000 / width))),
            max(1, min(10_000, round(bottom * 10_000 / height))),
        ]
        if bbox[0] >= bbox[2] or bbox[1] >= bbox[3]:
            raise ProcessingError("PP-OCR produced an empty normalized box")
        spans.append(
            {
                "text": text,
                "bbox": bbox,
                "confidenceMilli": max(0, min(1_000, round(float(raw_score) * 1_000))),
            }
        )
    text = "\n".join(span["text"] for span in spans)
    return page_text(page, "ppOcrV6Small", width_millipoints, height_millipoints, text, spans)


@lru_cache(maxsize=1)
def ocr_pipeline():
    device = processor_device()
    require_device(device)
    manifest = load_model_manifest()
    from paddleocr import PaddleOCR

    return PaddleOCR(
        text_detection_model_name=DET_MODEL,
        text_detection_model_dir=manifest["models"][DET_MODEL]["path"],
        text_recognition_model_name=REC_MODEL,
        text_recognition_model_dir=manifest["models"][REC_MODEL]["path"],
        use_doc_orientation_classify=False,
        use_doc_unwarping=False,
        use_textline_orientation=False,
        device=device,
        cpu_threads=max(1, min(8, os.cpu_count() or 1)),
        enable_mkldnn=False,
    )


def process_pdf(path: Path) -> Iterator[dict]:
    import pypdfium2 as pdfium

    try:
        document = pdfium.PdfDocument(str(path))
    except Exception as error:
        raise ProcessingError("PDFium could not open the PDF") from error
    with document:
        page_count = len(document)
        print(f"PDF pageCount={page_count}", file=sys.stderr, flush=True)
        if page_count < 1:
            raise ProcessingError(f"PDF contains no pages: pageCount={page_count}")
        with tempfile.TemporaryDirectory(prefix="centaeris-pdf-") as directory:
            render_root = Path(directory)
            for index in range(page_count):
                page_number = index + 1
                with closing(document[index]) as page:
                    width_points, height_points = page.get_size()
                    width_millipoints = max(1, round(width_points * 1_000))
                    height_millipoints = max(1, round(height_points * 1_000))
                    try:
                        with closing(page.get_textpage()) as text_page:
                            if text_page.count_chars() > MAX_PAGE_TEXT_BYTES:
                                raise ProcessingError(f"PDF page {page_number} text exceeds the character bound")
                            native = normalize_text(text_page.get_text_range())
                    except ProcessingError:
                        raise
                    except Exception as error:
                        raise ProcessingError(
                            f"PDFium text extraction failed on page {page_number}"
                        ) from error
                    if native_text_is_usable(native):
                        yield page_text(
                            page_number,
                            "pdfNative",
                            width_millipoints,
                            height_millipoints,
                            native,
                            [{"text": native, "bbox": [0, 0, 10_000, 10_000]}],
                        )
                        continue
                    pixel_width = max(1, round(width_points * RENDER_DPI / 72))
                    pixel_height = max(1, round(height_points * RENDER_DPI / 72))
                    if pixel_width * pixel_height > MAX_RENDERED_PIXELS_PER_PAGE:
                        raise ProcessingError(
                            f"PDF page {page_number} exceeds the rendered pixel limit"
                        )
                    try:
                        with closing(page.render(scale=RENDER_DPI / 72)) as bitmap:
                            image = bitmap.to_pil().convert("RGB")
                    except Exception as error:
                        raise ProcessingError(
                            f"PDFium rendering failed on page {page_number}"
                        ) from error
                    with image:
                        if image_is_blank(image):
                            result = empty_page(page_number, width_millipoints, height_millipoints)
                        else:
                            rendered = render_root / "page.png"
                            try:
                                image.save(rendered, format="PNG")
                                result = ocr_image(rendered, page_number, width_millipoints, height_millipoints)
                            finally:
                                rendered.unlink(missing_ok=True)
                    yield result


def process_image(path: Path) -> Iterator[dict]:
    from PIL import Image

    Image.MAX_IMAGE_PIXELS = MAX_RENDERED_PIXELS_PER_PAGE
    try:
        with Image.open(path) as source:
            frames = getattr(source, "n_frames", 1)
            print(f"image pageCount={frames}", file=sys.stderr, flush=True)
            if frames < 1:
                raise ProcessingError(f"image contains no frames: pageCount={frames}")
            with tempfile.TemporaryDirectory(prefix="centaeris-image-") as directory:
                root = Path(directory)
                for index in range(frames):
                    source.seek(index)
                    if source.width * source.height > MAX_RENDERED_PIXELS_PER_PAGE:
                        raise ProcessingError("image exceeds the pixel limit")
                    with source.convert("RGB") as image:
                        width = max(1, round(image.width * 72_000 / RENDER_DPI))
                        height = max(1, round(image.height * 72_000 / RENDER_DPI))
                        if image_is_blank(image):
                            result = empty_page(index + 1, width, height)
                        else:
                            rendered = root / "page.png"
                            try:
                                image.save(rendered, format="PNG")
                                result = ocr_image(rendered, index + 1, width, height)
                            finally:
                                rendered.unlink(missing_ok=True)
                    yield result
    except ProcessingError:
        raise
    except Exception as error:
        raise ProcessingError("image decoding failed") from error


def process_text(path: Path) -> Iterator[dict]:
    try:
        text = normalize_text(path.read_text(encoding="utf-8", errors="strict"))
    except (OSError, UnicodeDecodeError) as error:
        raise ProcessingError("text input is not valid UTF-8") from error
    if not text:
        yield empty_page(1, 1_000, 1_000)
        return
    yield page_text(
        1,
        "nativeText",
        1_000,
        1_000,
        text,
        [{"text": text, "bbox": [0, 0, 10_000, 10_000]}],
    )


def convert_office(path: Path, root: Path) -> Path:
    profile = root / "libreoffice-profile"
    output = root / "converted"
    profile.mkdir()
    output.mkdir()
    command = [
        "soffice",
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--headless",
        "--safe-mode",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output),
        str(path),
    ]
    try:
        result = subprocess.run(command, capture_output=True, timeout=300, check=False)
    except (OSError, subprocess.TimeoutExpired) as error:
        raise ProcessingError("LibreOffice conversion failed") from error
    candidates = list(output.glob("*.pdf"))
    if result.returncode != 0 or len(candidates) != 1:
        raise ProcessingError("LibreOffice conversion failed")
    return candidates[0]


def _xlsx_color(color) -> str | None:
    if color is None or color.type != "rgb" or not isinstance(color.rgb, str):
        return None
    value = color.rgb.upper()
    if len(value) == 8:
        value = value[2:]
    return f"#{value}" if len(value) == 6 and all(character in "0123456789ABCDEF" for character in value) else None


def _xlsx_border(side) -> dict | None:
    if side is None or not side.style:
        return None
    return {"style": side.style, "color": _xlsx_color(side.color)}


def _xlsx_style(cell) -> dict:
    borders = {
        name: value
        for name in ("top", "right", "bottom", "left")
        if (value := _xlsx_border(getattr(cell.border, name))) is not None
    }
    font_color = _xlsx_color(cell.font.color)
    fill = _xlsx_color(cell.fill.fgColor) if cell.fill.fill_type == "solid" else None
    return {
        "font": {"bold": bool(cell.font.bold), "italic": bool(cell.font.italic), "color": font_color},
        "fill": fill,
        "horizontal": cell.alignment.horizontal,
        "vertical": cell.alignment.vertical,
        "wrapText": bool(cell.alignment.wrap_text),
        "borders": borders,
    }


def _xlsx_cell_value(value) -> tuple[str, str]:
    if value is None:
        return "", "blank"
    if isinstance(value, datetime):
        if value.time() == datetime_time():
            return value.date().isoformat(), "date"
        return value.isoformat(timespec="seconds"), "date"
    if isinstance(value, date):
        return value.isoformat(), "date"
    if isinstance(value, datetime_time):
        return value.isoformat(timespec="seconds"), "date"
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), "boolean"
    if isinstance(value, (int, float)):
        return str(value), "number"
    return str(value), "text"


def workbook_preview(path: Path) -> dict | None:
    """Return a bounded, inert workbook model; formulas are never evaluated."""
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell

    try:
        workbook = load_workbook(path, read_only=False, data_only=False, keep_links=False)
        cached = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    except Exception as error:
        raise ProcessingError("spreadsheet workbook is invalid") from error
    if len(workbook.worksheets) > MAX_WORKBOOK_PREVIEW_SHEETS:
        workbook.close()
        cached.close()
        return None
    styles, style_ids, sheets, cell_count = [], {}, [], 0
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets):
            cached_sheet = cached.worksheets[sheet_index]
            max_row = max([sheet.max_row, *sheet.row_dimensions.keys()], default=1)
            max_column = max(sheet.max_column, 1)
            if sheet.column_dimensions:
                max_column = max(max_column, *(dimension.max for dimension in sheet.column_dimensions.values()))
            if max_row * max_column > MAX_WORKBOOK_PREVIEW_GRID_CELLS:
                return None
            rows = []
            for row_index in range(1, max_row + 1):
                dimension = sheet.row_dimensions.get(row_index)
                row = []
                for cell in sheet[row_index]:
                    if isinstance(cell, MergedCell) or cell.value is None:
                        continue
                    cell_count += 1
                    if cell_count > MAX_WORKBOOK_PREVIEW_CELLS:
                        return None
                    value, kind = _xlsx_cell_value(cell.value)
                    formula = value if cell.data_type == "f" else None
                    if formula is not None:
                        cached_value = cached_sheet.cell(cell.row, cell.column).value
                        value, kind = _xlsx_cell_value(cached_value if cached_value is not None else cell.value)
                        kind = "formula"
                    style = _xlsx_style(cell)
                    style_key = json.dumps(style, sort_keys=True, separators=(",", ":"))
                    if style_key not in style_ids:
                        style_ids[style_key] = len(styles)
                        styles.append(style)
                    row.append({
                        "column": cell.column,
                        "displayValue": value,
                        "kind": kind,
                        "formula": formula,
                        "numberFormat": cell.number_format,
                        "styleId": style_ids[style_key],
                    })
                if row or dimension is not None:
                    rows.append({
                        "index": row_index,
                        "heightPx": round(dimension.height * 96 / 72) if dimension and dimension.height else None,
                        "hidden": bool(dimension.hidden) if dimension else False,
                        "cells": row,
                    })
            columns = []
            for dimension in sheet.column_dimensions.values():
                for column_index in range(dimension.min, dimension.max + 1):
                    columns.append({
                        "index": column_index,
                        "widthPx": round(float(dimension.width) * 7 + 5) if dimension.width else None,
                        "hidden": bool(dimension.hidden),
                    })
            sheets.append({
                "name": sheet.title,
                "state": sheet.sheet_state,
                "maxRow": max_row,
                "maxColumn": max_column,
                "frozenPane": str(sheet.freeze_panes) if sheet.freeze_panes else None,
                "columns": columns,
                "rows": rows,
                "merges": [{
                    "startRow": merged.min_row,
                    "startColumn": merged.min_col,
                    "endRow": merged.max_row,
                    "endColumn": merged.max_col,
                } for merged in sheet.merged_cells.ranges],
            })
        result = {"schema": "knowledge.workbook_preview.v1", "styles": styles, "sheets": sheets}
        if len(json.dumps(result, ensure_ascii=False, separators=(",", ":")).encode("utf-8")) > MAX_WORKBOOK_PREVIEW_BYTES:
            return None
        return result
    finally:
        workbook.close()
        cached.close()


def write_bounded(output, content: bytes, maximum: int) -> None:
    if output.tell() + len(content) > maximum:
        raise ProcessingError(f"{Path(output.name).name} exceeds the output byte limit")
    output.write(content)


def write_outputs(display_name: str, pages: Iterator[dict], output: Path, preview: Path | None, workbook: dict | None = None) -> int:
    preview_size = preview.stat().st_size if preview else 0
    workbook_bytes = json.dumps(workbook, ensure_ascii=False, separators=(",", ":"), sort_keys=True).encode("utf-8") if workbook else b""
    if len(workbook_bytes) > MAX_WORKBOOK_PREVIEW_BYTES:
        raise ProcessingError("workbook.json exceeds the output byte limit")
    if preview_size + len(workbook_bytes) > MAX_OUTPUT_BYTES:
        raise ProcessingError("preview.pdf exceeds the output byte limit")
    canonical_limit = MAX_OUTPUT_BYTES - preview_size - len(workbook_bytes)
    encoder = json.JSONEncoder(ensure_ascii=False, separators=(",", ":"), sort_keys=True)
    # Stage every output until processing and all budgets succeed. Runtime only
    # commits after a successful process exit; failed pages never become ready.
    with tempfile.TemporaryDirectory(prefix=".processing-", dir=output) as directory:
        root = Path(directory)
        page_count = 0
        with (root / "canonical.md").open("wb") as canonical, (root / "manifest.json").open("wb") as manifest:
            heading = f"# {display_name}\n\n".encode("utf-8")
            write_bounded(canonical, heading, canonical_limit)
            line = heading.count(b"\n") + 1
            write_bounded(manifest, b'{"schema":"knowledge.derived_manifest.v1","pages":[', MAX_MANIFEST_BYTES)
            for page in pages:
                header = f"## Page {page['page']}\n\n".encode("utf-8")
                body = page["text"].encode("utf-8")
                if len(body) > MAX_PAGE_TEXT_BYTES:
                    raise ProcessingError(f"page {page['page']} text exceeds the byte limit")
                write_bounded(canonical, header, canonical_limit)
                start_byte = canonical.tell()
                start_line = line + header.count(b"\n")
                write_bounded(canonical, body, canonical_limit)
                end_byte = canonical.tell()
                end_line = start_line + body.count(b"\n")
                write_bounded(canonical, b"\n\n", canonical_limit)
                entry = {
                    "pageText": page,
                    "canonicalStartByte": start_byte,
                    "canonicalEndByte": end_byte,
                    "canonicalStartLine": start_line,
                    "canonicalEndLine": end_line,
                }
                if page_count:
                    write_bounded(manifest, b",", MAX_MANIFEST_BYTES)
                for chunk in encoder.iterencode(entry):
                    write_bounded(manifest, chunk.encode("utf-8"), MAX_MANIFEST_BYTES)
                page_count += 1
                line = end_line + 2
            if not page_count:
                raise ProcessingError("document contains no pages: pageCount=0")
            write_bounded(manifest, f'],"pageCount":{page_count}}}'.encode(), MAX_MANIFEST_BYTES)
            canonical_size = canonical.tell()
        if preview:
            with preview.open("rb") as source, (root / "preview.pdf").open("wb") as target:
                while chunk := source.read(64 * 1024):
                    write_bounded(target, chunk, MAX_OUTPUT_BYTES - canonical_size)
        if workbook_bytes:
            (root / "workbook.json").write_bytes(workbook_bytes)
        for name in ("canonical.md", "preview.pdf", "workbook.json", "manifest.json"):
            if (root / name).is_file():
                os.replace(root / name, output / name)
    return page_count


def process(request_path: Path) -> None:
    started = time.monotonic()
    try:
        request = json.loads(request_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise ProcessingError("processing request is invalid") from error
    if set(request) != {"schema", "inputPath", "displayName", "contentType", "outputDirectory"} or request.get(
        "schema"
    ) != "knowledge.processing.request.v1":
        raise ProcessingError("processing request fields are invalid")
    source = Path(request["inputPath"])
    output = Path(request["outputDirectory"])
    display_name = request["displayName"]
    if (
        not source.is_file()
        or source.stat().st_size > MAX_INPUT_BYTES
        or not isinstance(display_name, str)
        or not display_name.strip()
        or display_name != display_name.strip()
    ):
        raise ProcessingError("processing input is invalid")
    output.mkdir(parents=True, exist_ok=True)
    suffix = Path(display_name).suffix.lower()
    with tempfile.TemporaryDirectory(prefix="centaeris-process-") as directory:
        preview = None
        workbook = None
        if suffix in TEXT_SUFFIXES:
            pages = process_text(source)
        elif suffix == ".pdf" or request["contentType"] == "application/pdf":
            pages = process_pdf(source)
        elif suffix in IMAGE_SUFFIXES or str(request["contentType"]).startswith("image/"):
            pages = process_image(source)
        elif suffix in OFFICE_SUFFIXES:
            if suffix == ".xlsx":
                workbook = workbook_preview(source)
            preview = convert_office(source, Path(directory))
            pages = process_pdf(preview)
        else:
            raise ProcessingError(f"unsupported document type: {suffix or 'none'}")
        with closing(pages):
            page_count = write_outputs(display_name, pages, output, preview, workbook)
    print(f"document processing completed: pageCount={page_count}; elapsedMs={round((time.monotonic() - started) * 1000)}", file=sys.stderr, flush=True)


def main(argv: list[str]) -> int:
    try:
        if argv == ["spec"]:
            print(json.dumps(specification(), separators=(",", ":"), sort_keys=True))
        elif argv == ["fetch-models"]:
            fetch_models()
        elif len(argv) == 2 and argv[0] == "process":
            process(Path(argv[1]))
        else:
            raise ProcessingError("usage: document_processor.py spec|fetch-models|process REQUEST")
    except ProcessingError as error:
        print(str(error), file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
